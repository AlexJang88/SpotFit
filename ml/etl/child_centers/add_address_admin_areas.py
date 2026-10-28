from __future__ import annotations

import csv
import re
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
TARGET_FILE = ROOT / "data" / "interim" / "지역아동센터_2021_2025_통합.csv"
STANDARD_FILE = ROOT / "data" / "raw" / "1년치만 있는 데이터" / "전국지역아동센터표준데이터.csv"
CROSSWALK_FILE = ROOT / "data" / "raw" / "서울시 상권분석서비스" / "법정동코드-행정동코드.xlsx"

ADDRESS_API = "https://www.juso.go.kr/api/solr/solrKeywordSearch"
SEOUL = "서울특별시"
MAX_WORKERS = 8

ALIASES = {
    "자치구": ("자치구", "시군구명", "시군구"),
    "시설명": ("시설명", "센터명"),
    "주소": ("주소", "시설 소재지", "소재지도로명주소", "소재지지번주소", "소재지"),
    "전화번호": ("시설 전화번호", "대표 전화번호", "전화번호"),
}

DROP_FIELDS = {
    "연번",
    "시설 전화번호",
    "팩스 번호",
    "시도",
    "시군구",
    "팩스번호",
    "홈페이지",
    "대표 전화번호",
    "입소정원",
    "아동현원",
    "시설 소재지",
    "홈페이지 주소",
}

# 원본에서 주소가 공란/오표기된 센터만 공식 기관 페이지의 주소로 보완한다.
# - 서울SOS지역아동복지센터(문화체육관광부 문화품앗이):
#   https://csv.culture.go.kr/frt/biz/intro/bizEntView.do?selectedId=USRCNFRM_00000084380
# - 구립우등생지역아동센터(서울시 우리동네키움포털):
#   https://umppa.seoul.go.kr/icare/user/fcltyInfoManage/BD_selectFcltyInfoManage.do?q_fcltyId=GD200103&q_fclty=1003
OFFICIAL_CENTER_ADDRESSES = {
    ("양천구", "서울SOS지역아동복지센터"): "서울특별시 양천구 가로공원로58길 32",
    ("강동구", "구립우등생지역아동센터"): "서울특별시 강동구 양재대로 1458",
}

# 주소검색이 두 행정동을 함께 반환하는 경계 건물은 서울시 센터 좌표와
# 서울시 행정동 경계(프로젝트의 frontend/public/seoul-dong.geojson)로 단일 판정했다.
# 서울시 상세: https://umppa.seoul.go.kr/icare/user/fcltyInfoManage/BD_selectFcltyInfoManage.do?q_fcltyId=GR200155&q_fclty=1003
OFFICIAL_CENTER_DONG_OVERRIDES = {
    ("구로구", "메리워드지역아동센터"): "오류제1동",
}

_thread_local = threading.local()


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def first_value(record: dict[str, str], aliases: tuple[str, ...]) -> str:
    return next((record.get(alias, "") for alias in aliases if record.get(alias, "")), "")


def normalize_name(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]", "", value).lower()


def normalize_phone(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    return digits[-8:] if len(digits) >= 8 else digits


def extract_gu(record: dict[str, str]) -> str:
    source_gu = first_value(record, ALIASES["자치구"])
    if re.fullmatch(r"[가-힣]+구", source_gu):
        return source_gu

    address = first_value(record, ALIASES["주소"])
    matched = re.search(r"서울(?:특별시)?\s+([가-힣]+구)", address)
    return matched.group(1) if matched else ""


def is_seoul(record: dict[str, str]) -> bool:
    province = clean(record.get("시도명")) or clean(record.get("시도"))
    if province:
        return province in {"서울", SEOUL}
    return bool(re.search(r"서울(?:특별시)?\s+", first_value(record, ALIASES["주소"])))


def address_key(address: str, gu: str) -> tuple[str, str, int, int] | None:
    """Return an exact road/lot building key; reject values that only look like floor/unit numbers."""
    address = re.sub(r"\s+", " ", address).strip()

    road = re.search(r"([가-힣A-Za-z0-9.·]+(?:대로|로|길))\s+(\d+)(?:-(\d+))?", address)
    if road:
        suffix = address[road.end() :]
        if not road.group(3) and re.match(r"^\s*(?:층|호)(?:\s|$)", suffix):
            return None
        return "road", road.group(1), int(road.group(2)), int(road.group(3) or 0)

    lot = re.search(r"([가-힣0-9]+(?:동|가))\s+(?:산\s*)?(\d+)(?:번지)?(?:\s*[-－]\s*(\d+))?", address)
    if lot:
        return "lot", lot.group(1), int(lot.group(2)), int(lot.group(3) or 0)

    return None


def query_text(gu: str, key: tuple[str, str, int, int]) -> str:
    _, road_or_dong, main_no, sub_no = key
    number = str(main_no) if sub_no == 0 else f"{main_no}-{sub_no}"
    return f"{SEOUL} {gu} {road_or_dong} {number}"


def session() -> requests.Session:
    if not hasattr(_thread_local, "session"):
        current = requests.Session()
        current.headers.update(
            {
                "User-Agent": "Mozilla/5.0 (compatible; SpotFit-address-enrichment/1.0)",
                "Origin": "https://www.juso.go.kr",
                "Referer": "https://www.juso.go.kr/",
            }
        )
        _thread_local.session = current
    return _thread_local.session


def exact_admin_dong(gu: str, key: tuple[str, str, int, int]) -> str:
    kind, road_or_dong, main_no, sub_no = key
    payload: dict[str, Any] = {
        "keyword": query_text(gu, key),
        "strSearchType": "HSTRY",
        "strFirstSort": "none",
        "strAblYn": "N",
        "strAotYn": "N",
        "strSynnYn": None,
        "strHstryYn": "Y",
        "reqFrom": "RN_SEARCH_KOR_WEB",
        "checkMoblieYn": "N",
        "strFunctionName": "Y",
        "pageable": {"page": 0, "size": 20},
    }

    for attempt in range(3):
        try:
            response = session().post(ADDRESS_API, json=payload, timeout=25)
            response.raise_for_status()
            items = response.json().get("results", {}).get("content", [])
            matches: list[dict[str, Any]] = []
            for item in items:
                if clean(item.get("sggNm")) != gu:
                    continue
                if kind == "road":
                    exact = (
                        clean(item.get("rdNmStr")) == road_or_dong
                        and int(item.get("bdMaSn") or 0) == main_no
                        and int(item.get("bdSbSn") or 0) == sub_no
                    )
                else:
                    exact = (
                        clean(item.get("emdNm")) == road_or_dong
                        and int(item.get("lndnMaSn") or 0) == main_no
                        and int(item.get("lndnSbSn") or 0) == sub_no
                    )
                if exact:
                    matches.append(item)

            dongs = {clean(item.get("hemdNm")) for item in matches if clean(item.get("hemdNm"))}
            return next(iter(dongs)) if len(dongs) == 1 else ""
        except (requests.RequestException, ValueError, TypeError):
            if attempt == 2:
                return ""
            time.sleep(0.5 * (attempt + 1))
    return ""


def load_current_crosswalk() -> tuple[dict[tuple[str, str], set[str]], dict[str, set[str]]]:
    worksheet = load_workbook(CROSSWALK_FILE, read_only=True, data_only=True)["KIKmix"]
    rows = worksheet.iter_rows(values_only=True)
    next(rows)

    legal_to_admin: dict[tuple[str, str], set[str]] = defaultdict(set)
    admin_names: dict[str, set[str]] = defaultdict(set)
    for values in rows:
        admin_name = clean(values[3])
        legal_name = clean(values[5])
        if clean(values[1]) != SEOUL or not clean(values[2]) or not admin_name or not legal_name:
            continue
        if clean(values[7]):
            continue
        gu = clean(values[2])
        legal_to_admin[(gu, legal_name)].add(admin_name)
        admin_names[gu].add(admin_name)
    return legal_to_admin, admin_names


def address_tokens(address: str) -> list[str]:
    tokens: list[str] = []
    for parenthetical in re.findall(r"\(([^)]*)\)", address):
        tokens.extend(re.findall(r"[가-힣0-9]+(?:동|가)", parenthetical))
    if not tokens:
        lot = re.search(r"\s([가-힣0-9]+(?:동|가))\s+(?:산\s*)?\d", address)
        if lot:
            tokens.append(lot.group(1))
    return tokens


def crosswalk_admin_dong(
    gu: str,
    address: str,
    legal_to_admin: dict[tuple[str, str], set[str]],
    admin_names: dict[str, set[str]],
) -> str:
    candidates: set[str] = set()
    for token in address_tokens(address):
        if token in admin_names.get(gu, set()):
            candidates.add(token)
        mapped = legal_to_admin.get((gu, token), set())
        if len(mapped) == 1:
            candidates.update(mapped)
    return next(iter(candidates)) if len(candidates) == 1 else ""


def canonical_admin_dong(gu: str, dong: str, admin_names: dict[str, set[str]]) -> str:
    """Normalize an official API spelling to the active code-table spelling."""
    if dong in admin_names.get(gu, set()):
        return dong
    candidate = re.sub(r"^(.+?)(\d+동)$", r"\1제\2", dong)
    return candidate if candidate in admin_names.get(gu, set()) else dong


def reference_center_queries() -> dict[tuple[str, str], set[tuple[str, str, int, int]]]:
    result: dict[tuple[str, str], set[tuple[str, str, int, int]]] = defaultdict(set)
    with STANDARD_FILE.open(encoding="cp949", newline="") as handle:
        for row in csv.DictReader(handle):
            if row.get("시도명") != SEOUL:
                continue
            gu = clean(row.get("시군구명"))
            name = normalize_name(clean(row.get("센터명")))
            if not gu or not name:
                continue
            for field in ("소재지도로명주소", "소재지지번주소"):
                key = address_key(clean(row.get(field)), gu)
                if key:
                    result[(gu, name)].add(key)

    for (gu, center_name), address in OFFICIAL_CENTER_ADDRESSES.items():
        key = address_key(address, gu)
        if key:
            result[(gu, normalize_name(center_name))].add(key)
    return result


def center_key(record: dict[str, str], gu: str, include_phone: bool) -> tuple[str, ...]:
    name = normalize_name(first_value(record, ALIASES["시설명"]))
    if include_phone:
        return gu, name, normalize_phone(first_value(record, ALIASES["전화번호"]))
    return gu, name


def main() -> None:
    with TARGET_FILE.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        original_fields = reader.fieldnames or []
        all_records = list(reader)

    records = [record for record in all_records if is_seoul(record)]

    if not records:
        raise ValueError(f"빈 파일입니다: {TARGET_FILE}")

    standard_queries = reference_center_queries()
    row_keys: list[tuple[str, str, int, int] | None] = []
    lookup_pairs: set[tuple[str, tuple[str, str, int, int]]] = set()
    for record in records:
        gu = extract_gu(record)
        address = first_value(record, ALIASES["주소"])
        key = address_key(address, gu)
        row_keys.append(key)
        if gu and key:
            lookup_pairs.add((gu, key))

    for (gu, _), keys in standard_queries.items():
        lookup_pairs.update((gu, key) for key in keys)

    official_lookup: dict[tuple[str, tuple[str, str, int, int]], str] = {}
    completed = 0
    total = len(lookup_pairs)
    print(f"행정안전부 공식 주소검색: {total:,}개 고유 주소 조회 시작", flush=True)
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {executor.submit(exact_admin_dong, gu, key): (gu, key) for gu, key in lookup_pairs}
        for future in as_completed(futures):
            pair = futures[future]
            official_lookup[pair] = future.result()
            completed += 1
            if completed % 50 == 0 or completed == total:
                print(f"  조회 {completed:,}/{total:,}", flush=True)

    # 동시 조회 중 공식 검색 서버가 간헐적으로 빈 응답을 줄 수 있어 한 번 더 천천히 확인한다.
    retry_pairs = [pair for pair, dong in official_lookup.items() if not dong]
    if retry_pairs:
        print(f"  빈 응답 재확인: {len(retry_pairs):,}개", flush=True)
        with ThreadPoolExecutor(max_workers=4) as executor:
            futures = {executor.submit(exact_admin_dong, gu, key): (gu, key) for gu, key in retry_pairs}
            for future in as_completed(futures):
                pair = futures[future]
                retried = future.result()
                if retried:
                    official_lookup[pair] = retried

    legal_to_admin, admin_names = load_current_crosswalk()
    methods: Counter[str] = Counter()
    assigned: list[str] = []
    for record, key in zip(records, row_keys):
        gu = extract_gu(record)
        address = first_value(record, ALIASES["주소"])
        center = (gu, normalize_name(first_value(record, ALIASES["시설명"])))
        dong = OFFICIAL_CENTER_DONG_OVERRIDES.get(center, "")
        if dong:
            methods["공식좌표_행정동경계"] += 1
        else:
            dong = official_lookup.get((gu, key), "") if gu and key else ""
        if dong:
            dong = canonical_admin_dong(gu, dong, admin_names)
        if dong and center not in OFFICIAL_CENTER_DONG_OVERRIDES:
            methods["공식주소검색"] += 1
        elif not dong:
            dong = crosswalk_admin_dong(gu, address, legal_to_admin, admin_names)
            if dong:
                methods["공식행정동매핑"] += 1
        assigned.append(dong)

    # 공식 표준데이터/기관 페이지가 원본의 공란·오표기 행에 정상 주소를 보완한다.
    standard_dongs: dict[tuple[str, str], set[str]] = defaultdict(set)
    for center, keys in standard_queries.items():
        gu, _ = center
        for key in keys:
            dong = official_lookup.get((gu, key), "")
            if dong:
                standard_dongs[center].add(dong)

    for index, record in enumerate(records):
        if assigned[index]:
            continue
        gu = extract_gu(record)
        center = (gu, normalize_name(first_value(record, ALIASES["시설명"])))
        dongs = standard_dongs.get(center, set())
        if len(dongs) == 1:
            assigned[index] = next(iter(dongs))
            methods["공식센터주소보완"] += 1

    # Repeated yearly rows may contain a blank/truncated address in one year only.
    by_phone: dict[tuple[str, ...], set[str]] = defaultdict(set)
    by_name: dict[tuple[str, ...], set[str]] = defaultdict(set)
    for record, dong in zip(records, assigned):
        if not dong:
            continue
        gu = extract_gu(record)
        by_phone[center_key(record, gu, True)].add(dong)
        by_name[center_key(record, gu, False)].add(dong)

    for index, record in enumerate(records):
        if assigned[index]:
            continue
        gu = extract_gu(record)
        phone_values = by_phone.get(center_key(record, gu, True), set())
        name_values = by_name.get(center_key(record, gu, False), set())
        if len(phone_values) == 1:
            assigned[index] = next(iter(phone_values))
            methods["동일센터_전화번호"] += 1
        elif len(name_values) == 1:
            assigned[index] = next(iter(name_values))
            methods["동일센터_시설명"] += 1

    missing_gu = [index + 2 for index, record in enumerate(records) if not extract_gu(record)]
    missing_dong = [index + 2 for index, dong in enumerate(assigned) if not dong]
    if missing_gu or missing_dong:
        print(f"자치구 미확인 행: {missing_gu}")
        print(f"행정동 미확인 행: {missing_dong}")
        raise ValueError("공식 근거로 확인되지 않은 행이 있어 결과 파일을 덮어쓰지 않았습니다.")

    for record, dong in zip(records, assigned):
        record["자치구"] = extract_gu(record)
        record["행정동"] = dong

    output_fields = [
        field for field in original_fields if field not in DROP_FIELDS | {"자치구", "행정동"}
    ]
    insert_at = output_fields.index("기준연도") + 1 if "기준연도" in output_fields else 0
    output_fields[insert_at:insert_at] = ["자치구", "행정동"]

    temporary = TARGET_FILE.with_suffix(".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=output_fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
    temporary.replace(TARGET_FILE)

    print(
        f"완료: 전체 {len(all_records):,}행 중 서울 {len(records):,}행 / "
        "자치구·행정동 빈칸 0건",
        flush=True,
    )
    print("판정 근거: " + ", ".join(f"{key} {value:,}행" for key, value in methods.items()), flush=True)


if __name__ == "__main__":
    main()
