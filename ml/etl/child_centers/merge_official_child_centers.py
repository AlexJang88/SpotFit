from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE_DIR = ROOT / "data" / "raw" / "지역아동센터"
OUTPUT_DIR = ROOT / "data" / "processed" / "지역아동센터_공식원본_2021_2025_통합"

def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_header(rows: list[tuple[object, ...]]) -> int:
    for index, row in enumerate(rows):
        values = {clean(value).replace("\n", " ") for value in row}
        if "연번" in values and ("시설명" in values or "센터명" in values):
            return index
    raise ValueError("머리글 행을 찾지 못했습니다.")


def unique_headers(row: tuple[object, ...]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(row, start=1):
        base = clean(value).replace("\n", " ") or f"이름없는열_{index}"
        seen[base] = seen.get(base, 0) + 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    while headers and headers[-1].startswith("이름없는열_"):
        headers.pop()
    return headers


def main() -> None:
    source_files = sorted(SOURCE_DIR.glob("20??_*.xlsx"))
    if [int(path.name[:4]) for path in source_files] != list(range(2021, 2026)):
        raise FileNotFoundError("2021~2025년 공식 XLSX가 각각 1개씩 필요합니다.")

    all_records: list[dict[str, str | int]] = []
    original_columns: list[str] = []
    counts: dict[int, int] = {}

    for path in source_files:
        year = int(path.name[:4])
        worksheet = load_workbook(path, read_only=True, data_only=True).active
        rows = list(worksheet.iter_rows(values_only=True))
        header_index = find_header(rows)
        headers = unique_headers(rows[header_index])
        for header in headers:
            if header not in original_columns:
                original_columns.append(header)

        year_count = 0
        for source_row, values in enumerate(rows[header_index + 1 :], start=header_index + 2):
            record = {header: clean(values[index] if index < len(values) else None) for index, header in enumerate(headers)}
            facility_name = record.get("시설명", record.get("센터명", ""))
            if not facility_name:
                continue
            record.update({"기준연도": year, "원본파일": path.name})
            all_records.append(record)
            year_count += 1
        counts[year] = year_count

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    provenance_fields = ["기준연도", "원본파일"]
    union_fields = provenance_fields + [column for column in original_columns if column not in provenance_fields]
    output_path = OUTPUT_DIR / "지역아동센터_2021_2025_통합.csv"
    with output_path.open(
        "w", encoding="utf-8-sig", newline=""
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=union_fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(all_records)

    print(f"통합 완료: {len(all_records):,}행, {output_path}")
    for year, count in counts.items():
        print(f"{year}: {count:,}행")


if __name__ == "__main__":
    main()
