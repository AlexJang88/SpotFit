from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
SOURCE_ROOT = ROOT / "data" / "raw" / "학교" / "학교현황"
OUTPUT_FILE = ROOT / "data" / "interim" / "학교현황_2024_2026_통합.xlsx"

OUTPUT_SHEETS = [
    "상반기_총개황",
    "상반기_유치원",
    "상반기_초등학교",
    "상반기_중학교",
    "상반기_고등학교",
    "하반기_학교수비교",
    "하반기_주요현황비교",
    "하반기_학년별학생수",
]


def clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalized_header(value: str) -> str:
    value = re.sub(r"20\d{2}년\s*", "", value)
    return clean(value).strip("_")


def expanded_header_grid(worksheet, start_row: int, end_row: int) -> list[list[str]]:
    grid = [
        [clean(worksheet.cell(row, column).value) for column in range(1, worksheet.max_column + 1)]
        for row in range(start_row, end_row + 1)
    ]
    for merged in worksheet.merged_cells.ranges:
        if merged.max_row < start_row or merged.min_row > end_row:
            continue
        value = clean(worksheet.cell(merged.min_row, merged.min_col).value)
        for row in range(max(merged.min_row, start_row), min(merged.max_row, end_row) + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                grid[row - start_row][column - 1] = value
    return grid


def flat_headers(worksheet, start_row: int, end_row: int) -> list[str]:
    grid = expanded_header_grid(worksheet, start_row, end_row)
    headers: list[str] = []
    seen: Counter[str] = Counter()
    for column in range(worksheet.max_column):
        levels: list[str] = []
        for row in grid:
            value = normalized_header(row[column])
            if value and (not levels or value != levels[-1]):
                levels.append(value)
        base = "__".join(levels) if levels else f"이름없는열_{column + 1}"
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def table_spec(period: str, sheet_index: int, year: int, max_row: int) -> tuple[str, int, int, int]:
    if period == "상반기":
        header_ranges = [(3, 5), (3, 4), (3, 5), (3, 6), (3, 5)]
        data_starts = [6, 5, 6, 7, 6]
        return OUTPUT_SHEETS[sheet_index], *header_ranges[sheet_index], data_starts[sheet_index]

    if sheet_index == 0:
        return OUTPUT_SHEETS[5], 6, 8, 9
    if sheet_index == 1:
        return (OUTPUT_SHEETS[6], 7, 9, 10) if year == 2024 else (OUTPUT_SHEETS[6], 8, 10, 11)
    if sheet_index == 2:
        return (OUTPUT_SHEETS[7], 5, 5, 6) if year == 2024 else (OUTPUT_SHEETS[7], 6, 6, 7)
    raise ValueError(f"예상하지 못한 하반기 시트 번호: {sheet_index}")


def main() -> None:
    source_files = sorted(SOURCE_ROOT.rglob("*.xlsx"))
    if len(source_files) != 4:
        raise FileNotFoundError(f"학교현황 XLSX 4개가 필요합니다. 현재 {len(source_files)}개")

    fields_by_sheet: dict[str, list[str]] = defaultdict(lambda: ["기준연도", "기준시점", "원본파일", "원본시트"])
    records_by_sheet: dict[str, list[dict[str, object]]] = defaultdict(list)

    for path in source_files:
        period = path.parent.name
        year_match = re.search(r"20\d{2}", path.name)
        if not year_match:
            raise ValueError(f"파일명에서 연도를 찾지 못했습니다: {path.name}")
        year = int(year_match.group())
        workbook = load_workbook(path, read_only=False, data_only=True)

        for sheet_index, worksheet in enumerate(workbook.worksheets):
            output_sheet, header_start, header_end, data_start = table_spec(
                period, sheet_index, year, worksheet.max_row
            )
            headers = flat_headers(worksheet, header_start, header_end)
            for header in headers:
                if header not in fields_by_sheet[output_sheet]:
                    fields_by_sheet[output_sheet].append(header)

            for values in worksheet.iter_rows(min_row=data_start, max_row=worksheet.max_row, values_only=True):
                if not any(clean(value) for value in values):
                    continue
                if clean(values[0] if values else None).startswith("주 "):
                    continue
                record: dict[str, object] = {
                    "기준연도": year,
                    "기준시점": period,
                    "원본파일": path.name,
                    "원본시트": worksheet.title,
                }
                record.update(
                    {header: values[index] if index < len(values) else None for index, header in enumerate(headers)}
                )
                records_by_sheet[output_sheet].append(record)

    workbook = Workbook()
    guide = workbook.active
    guide.title = "안내"
    guide_rows = [
        ["항목", "내용"],
        ["통합 대상", "data/raw/학교/학교현황의 공식 XLSX 4개"],
        ["상반기", "2025·2026년 4월 1일 기준 현황"],
        ["하반기", "2024·2025년 상·하반기 비교표"],
        ["통합 방식", "같은 통계표끼리 연도 방향으로 결합"],
        ["주의", "상반기 현황표와 하반기 비교표는 집계 기준이 달라 서로 직접 합산하지 않음"],
        ["추가 열", "기준연도, 기준시점, 원본파일, 원본시트"],
    ]
    for row in guide_rows:
        guide.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for output_sheet in OUTPUT_SHEETS:
        worksheet = workbook.create_sheet(output_sheet)
        fields = fields_by_sheet[output_sheet]
        worksheet.append(fields)
        for record in records_by_sheet[output_sheet]:
            worksheet.append([record.get(field) for field in fields])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for column in range(1, worksheet.max_column + 1):
            max_length = max(len(clean(worksheet.cell(row, column).value)) for row in range(1, worksheet.max_row + 1))
            worksheet.column_dimensions[get_column_letter(column)].width = min(max(max_length + 2, 10), 45)

    for cell in guide[1]:
        cell.fill = header_fill
        cell.font = header_font
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 90
    guide.freeze_panes = "A2"

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_FILE.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(OUTPUT_FILE)

    print(f"통합 완료: {OUTPUT_FILE}")
    for sheet in OUTPUT_SHEETS:
        print(f"{sheet}: {len(records_by_sheet[sheet]):,}행, {len(fields_by_sheet[sheet]):,}열")


if __name__ == "__main__":
    main()
