from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from ml.etl.child_centers.add_address_admin_areas import (
    address_key,
    canonical_admin_dong,
    crosswalk_admin_dong,
    exact_admin_dong,
    load_current_crosswalk,
)


SOURCE_FILE = ROOT / "data" / "interim" / "학교별_일람표_2021_2026_상하반기_통합.csv"
OUTPUT_FILE = ROOT / "data" / "interim" / "유치원_개별현황_2021_2026_상하반기.xlsx"
FINAL_FIELDS = {
    "기준연도": ("기준연도",),
    "기준시점": ("기준시점",),
    "원본파일": ("원본파일",),
    "조사기준일": ("학교 기본 정보__조사기준일",),
    "KEDI학교코드": (
        "학교 기본 정보__KEDI학교코드",
        "학교 기본 정보__학교코드 (KEDI)",
    ),
    "학교명": ("학교 기본 정보__학교명",),
    "주소": ("학교 기본 정보__주소",),
    "자치구": ("학교 기본 정보__행정구",),
    "설립구분": ("학교 기본 정보__설립",),
    "학교상태": ("학교 기본 정보__상태",),
    "학급수": ("학년별 편성 학급수 (개)__편성학급수_계",),
    "원아수_계": ("학년별 학생수 (명)__총계__학생수_총계_계",),
    "특수학급수": (
        "학급 구분별 학급수 및 학생수__특수학급(편성학급)__특수학급_학급수",
    ),
    "특수원아수_계": (
        "학급 구분별 학급수 및 학생수__특수학급(편성학급)__특수학급_학생수_계",
    ),
    "교원수_계": (
        "교원수 (명)__총계__교원수_총계_계",
        "교원수 (명)__총계__교원수_총계__교원수_총계_계",
    ),
    "특수교사수_계": (
        "교원수 (명)__특수교사(정규)__교원수_정규_특수_계",
        "교원수 (명)__정규교원__교원수_정규교원_특수교사__교원수_정규_특수_계",
    ),
    "학급당_원아수": ("학급당 학생수 (명)__학급당 학생수",),
    "교원1인당_원아수": ("교원1인당 학생수 (명)__교원1인당 학생수",),
}

OUTPUT_FIELDS = list(FINAL_FIELDS)
OUTPUT_FIELDS.insert(OUTPUT_FIELDS.index("자치구") + 1, "행정동")
NUMBER_FIELDS = {
    "기준연도",
    "학급수",
    "원아수_계",
    "특수학급수",
    "특수원아수_계",
    "교원수_계",
    "특수교사수_계",
    "학급당_원아수",
    "교원1인당_원아수",
}

# 원본 주소가 폐지·오기되어 건물번호 완전일치 검색이 되지 않는 학교만 공식 자료로 보완한다.
# 근거: 행정안전부 주소검색의 정정 주소·동일 도로 구간, 교육부 유치원알리미,
# 서울시/자치구 공개자료 및 중앙선관위 투표소 주소.
OFFICIAL_DONG_OVERRIDES = {
    ("중랑구", "동국유치원"): "면목제4동",
    ("서대문구", "인왕유치원"): "홍제2동",
    ("은평구", "충암유치원"): "응암제3동",
    ("금천구", "럭키유치원"): "시흥제1동",
    ("금천구", "백상유치원"): "시흥제5동",
    ("강서구", "문성유치원"): "우장산동",
    ("양천구", "대한유치원"): "신월2동",
    ("강남구", "강남유치원"): "논현2동",
    ("강남구", "신양유치원"): "개포2동",
    ("관악구", "종인유치원"): "행운동",
    ("관악구", "나랑유치원"): "은천동",
    ("성북구", "성신유치원"): "동선동",
}


def coalesce(row: list[str], indexes: list[int], field_name: str) -> str:
    values = [row[index].strip() for index in indexes if row[index].strip()]
    distinct = list(dict.fromkeys(values))
    if len(distinct) > 1:
        raise ValueError(f"{field_name}에 서로 다른 원본값이 중복 존재합니다: {distinct}")
    return distinct[0] if distinct else ""


def normalize(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]", "", value).lower()


def address_gu(address: str) -> str:
    matched = re.search(r"서울(?:특별시)?\s+([가-힣]+구)", address)
    return matched.group(1) if matched else ""


def excel_value(field: str, value: str) -> object:
    if field not in NUMBER_FIELDS or not value:
        return value
    try:
        number = Decimal(value.replace(",", ""))
    except InvalidOperation:
        return value
    return int(number) if number == number.to_integral_value() else float(number)


def optimize_worksheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 24

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "기준연도": 10,
        "기준시점": 10,
        "원본파일": 34,
        "조사기준일": 13,
        "KEDI학교코드": 16,
        "학교명": 28,
        "주소": 62,
        "자치구": 12,
        "행정동": 14,
        "설립구분": 11,
        "학교상태": 13,
    }
    for index, field in enumerate(OUTPUT_FIELDS, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = widths.get(field, 16)


def add_admin_dongs(records: list[dict[str, str]]) -> None:
    legal_to_admin, admin_names = load_current_crosswalk()
    pairs: set[tuple[str, tuple[str, str, int, int]]] = set()
    keys: list[tuple[str, str, int, int] | None] = []
    for record in records:
        record["자치구"] = address_gu(record["주소"]) or record["자치구"]
        key = address_key(record["주소"], record["자치구"])
        keys.append(key)
        if key:
            pairs.add((record["자치구"], key))

    print(f"행정안전부 공식 주소검색: 고유 주소 {len(pairs):,}개", flush=True)
    official: dict[tuple[str, tuple[str, str, int, int]], str] = {}
    with ThreadPoolExecutor(max_workers=8) as executor:
        futures = {
            executor.submit(exact_admin_dong, gu, key): (gu, key) for gu, key in pairs
        }
        for completed, future in enumerate(as_completed(futures), 1):
            pair = futures[future]
            official[pair] = future.result()
            if completed % 100 == 0 or completed == len(futures):
                print(f"  조회 {completed:,}/{len(futures):,}", flush=True)

    for record, key in zip(records, keys):
        gu = record["자치구"]
        dong = OFFICIAL_DONG_OVERRIDES.get((gu, record["학교명"]), "")
        if not dong:
            dong = official.get((gu, key), "") if key else ""
        if dong:
            dong = canonical_admin_dong(gu, dong, admin_names)
        else:
            dong = crosswalk_admin_dong(
                gu, record["주소"], legal_to_admin, admin_names
            )
        record["행정동"] = dong

    # 한 시점의 주소 표기가 불완전한 경우, 공식 검색으로 확인된 동일 학교의 값을 보완한다.
    by_kedi: dict[str, set[str]] = defaultdict(set)
    by_name: dict[tuple[str, str], set[str]] = defaultdict(set)
    for record in records:
        if not record["행정동"]:
            continue
        if record["KEDI학교코드"]:
            by_kedi[record["KEDI학교코드"]].add(record["행정동"])
        by_name[(record["자치구"], normalize(record["학교명"]))].add(record["행정동"])

    for record in records:
        if record["행정동"]:
            continue
        kedi_dongs = by_kedi.get(record["KEDI학교코드"], set()) if record["KEDI학교코드"] else set()
        name_dongs = by_name.get((record["자치구"], normalize(record["학교명"])), set())
        if len(kedi_dongs) == 1:
            record["행정동"] = next(iter(kedi_dongs))
        elif len(name_dongs) == 1:
            record["행정동"] = next(iter(name_dongs))


def duplicate_count(records: list[dict[str, str]]) -> tuple[int, int]:
    kedi_seen: set[tuple[str, str]] = set()
    identity_seen: set[tuple[str, str, str]] = set()
    kedi_duplicates = 0
    identity_duplicates = 0
    for record in records:
        year = record["기준연도"]
        if record["KEDI학교코드"]:
            key = year, record["KEDI학교코드"]
            kedi_duplicates += key in kedi_seen
            kedi_seen.add(key)
        identity = year, normalize(record["학교명"]), normalize(record["주소"])
        identity_duplicates += identity in identity_seen
        identity_seen.add(identity)
    return kedi_duplicates, identity_duplicates


def main() -> None:
    with SOURCE_FILE.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        headers = next(reader)
        grade_index = next(
            index for index, header in enumerate(headers) if header.split("__")[-1] == "학교급"
        )

        records = [row for row in reader if row[grade_index].strip() == "유치원"]

    header_indexes = {header: index for index, header in enumerate(headers)}
    missing_fields = {
        source
        for sources in FINAL_FIELDS.values()
        for source in sources
        if source not in header_indexes
    }
    if missing_fields:
        raise KeyError(f"원본에 필요한 컬럼이 없습니다: {sorted(missing_fields)}")

    field_indexes = {
        output: [header_indexes[source] for source in sources]
        for output, sources in FINAL_FIELDS.items()
    }

    output_records: list[dict[str, str]] = []
    for row in records:
        output_records.append(
            {
                output: coalesce(row, indexes, output)
                for output, indexes in field_indexes.items()
            }
        )

    add_admin_dongs(output_records)
    upper = [record for record in output_records if record["기준시점"] == "상반기"]
    lower = [record for record in output_records if record["기준시점"] == "하반기"]
    sort_key = lambda record: (
        int(record["기준연도"]),
        record["자치구"],
        record["행정동"],
        record["학교명"],
    )
    upper.sort(key=sort_key)
    lower.sort(key=sort_key)
    upper_duplicates = duplicate_count(upper)
    lower_duplicates = duplicate_count(lower)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, sheet_records in (("상반기_분석", upper), ("하반기_검증", lower)):
        worksheet = workbook.create_sheet(title)
        worksheet.append(OUTPUT_FIELDS)
        for record in sheet_records:
            worksheet.append(
                [excel_value(field, record.get(field, "")) for field in OUTPUT_FIELDS]
            )
        optimize_worksheet(worksheet)

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_FILE.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(OUTPUT_FILE)
    missing_upper = sum(not record["행정동"] for record in upper)
    missing_lower = sum(not record["행정동"] for record in lower)
    print(
        f"완료: 상반기 {len(upper):,}행 / 하반기 {len(lower):,}행 / "
        f"{len(OUTPUT_FIELDS):,}열, {OUTPUT_FILE}"
    )
    print(f"행정동 미확인: 상반기 {missing_upper:,}행 / 하반기 {missing_lower:,}행")
    print(
        "연도 내 중복(KEDI / 학교명+주소): "
        f"상반기 {upper_duplicates[0]:,}/{upper_duplicates[1]:,}건, "
        f"하반기 {lower_duplicates[0]:,}/{lower_duplicates[1]:,}건"
    )
    print("최종 컬럼:")
    for field in OUTPUT_FIELDS:
        print(f"  {field}")


if __name__ == "__main__":
    main()
