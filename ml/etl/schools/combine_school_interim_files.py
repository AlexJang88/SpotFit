from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[3]
INTERIM_DIR = ROOT / "data" / "interim"
ROSTER_FILE = INTERIM_DIR / "학교별_일람표_2021_2026_상하반기_통합.csv"
STATUS_FILE = INTERIM_DIR / "학교현황_2024_2026_통합.xlsx"
OUTPUT_FILE = INTERIM_DIR / "학교데이터_통합.xlsx"


def main() -> None:
    if not ROSTER_FILE.exists() or not STATUS_FILE.exists():
        raise FileNotFoundError("통합할 학교별 일람표 CSV와 학교현황 XLSX가 모두 필요합니다.")

    output = Workbook(write_only=True)

    roster_sheet = output.create_sheet("학교별_일람표")
    roster_count = 0
    with ROSTER_FILE.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            roster_sheet.append(row)
            roster_count += 1

    status = load_workbook(STATUS_FILE, read_only=True, data_only=True)
    status_counts: list[tuple[str, int]] = []
    for source_sheet in status.worksheets:
        target_sheet = output.create_sheet(source_sheet.title)
        row_count = 0
        for row in source_sheet.iter_rows(values_only=True):
            target_sheet.append(list(row))
            row_count += 1
        status_counts.append((source_sheet.title, row_count))

    temporary = OUTPUT_FILE.with_suffix(".tmp.xlsx")
    output.save(temporary)
    temporary.replace(OUTPUT_FILE)

    print(f"완료: {OUTPUT_FILE}")
    print(f"학교별_일람표: {roster_count:,}행(헤더 포함)")
    for sheet_name, row_count in status_counts:
        print(f"{sheet_name}: {row_count:,}행(헤더 포함)")


if __name__ == "__main__":
    main()
