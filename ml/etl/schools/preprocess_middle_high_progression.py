from __future__ import annotations

import csv
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
SOURCE_FILE = ROOT / "data" / "raw" / "학교" / "고등학교_진학률_2011_2025.csv"
OUTPUT_FILE = ROOT / "data" / "interim" / "고등학교_진학률_2011_2025_1차전처리.xlsx"
SCHOOL_TYPES = ("전체", "일반고", "특성화고", "특수목적고", "자율고", "기타")
OUTPUT_FIELDS = [
    "연도", "지역구분", "지역", "고등학교_유형", "중학교_졸업자수",
    "고등학교_진학자수", "졸업자대비_고등학교진학률_pct",
    "미진학자수", "유형별_진학구성비_pct",
    "졸업자수_전년대비증감률_pct", "진학구성비_전년대비증감_pp",
]
COLUMN_LABELS = {
    "연도": "기준연도",
    "지역구분": "지역구분",
    "지역": "지역명",
    "고등학교_유형": "진학_고등학교유형",
    "중학교_졸업자수": "중학교_졸업자수(명)",
    "고등학교_진학자수": "해당유형_진학자수(명)",
    "졸업자대비_고등학교진학률_pct": "졸업자중_해당유형진학률(%)",
    "미진학자수": "고등학교_미진학자수(명)",
    "유형별_진학구성비_pct": "진학자중_해당유형비율(%)",
    "졸업자수_전년대비증감률_pct": "중학교졸업자수_전년대비증감률(%)",
    "진학구성비_전년대비증감_pp": "해당유형진학비율_전년대비증감(%p)",
}


def parse_integer(value: str, field: str) -> int | None:
    cleaned = value.strip().replace(",", "")
    if not cleaned:
        return None
    try:
        number = Decimal(cleaned)
    except InvalidOperation as error:
        raise ValueError(f"{field}이 숫자가 아닙니다: {value!r}") from error
    if number < 0 or number != number.to_integral_value():
        raise ValueError(f"{field}이 0 이상의 정수가 아닙니다: {value!r}")
    return int(number)


def parse_float(value: str, field: str) -> float | None:
    cleaned = value.strip().replace(",", "")
    if not cleaned:
        return None
    try:
        return float(Decimal(cleaned))
    except InvalidOperation as error:
        raise ValueError(f"{field}이 숫자가 아닙니다: {value!r}") from error


def percent_change(current: int, previous: int) -> float | None:
    if previous == 0:
        return None
    return round((current - previous) / previous * 100, 2)


def style_sheet(worksheet, widths: dict[str, int]) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = True
    worksheet.row_dimensions[1].height = 30
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, cell in enumerate(worksheet[1], 1):
        worksheet.column_dimensions[get_column_letter(index)].width = widths.get(str(cell.value), 18)


def main() -> None:
    with SOURCE_FILE.open(encoding="utf-8-sig", newline="") as handle:
        source_records = list(csv.DictReader(handle))

    records: list[dict[str, object]] = []
    for source in source_records:
        region = source["지역"].strip()
        records.append({
            "연도": int(source["연도"]),
            "지역구분": "서울전체" if region == "서울특별시" else "자치구",
            "지역": region,
            "고등학교_유형": source["고등학교_유형"].strip(),
            "중학교_졸업자수": parse_integer(source["중학교_졸업자수"], "중학교_졸업자수"),
            "고등학교_진학자수": parse_integer(source["고등학교_진학자수"], "고등학교_진학자수"),
            "졸업자대비_고등학교진학률_pct": parse_float(source["진학률_pct"], "진학률_pct"),
        })

    key_counts = Counter(
        (record["연도"], record["지역"], record["고등학교_유형"]) for record in records
    )
    duplicate_count = sum(count - 1 for count in key_counts.values() if count > 1)
    if duplicate_count:
        raise ValueError(f"연도+지역+유형 중복이 {duplicate_count}건 있습니다.")

    grouped: dict[tuple[int, str], list[dict[str, object]]] = defaultdict(list)
    for record in records:
        grouped[(int(record["연도"]), str(record["지역"]))].append(record)

    incomplete_type_groups = 0
    inconsistent_graduate_groups = 0
    detail_total_mismatches = 0
    for group_records in grouped.values():
        types = {str(record["고등학교_유형"]) for record in group_records}
        if types != set(SCHOOL_TYPES):
            incomplete_type_groups += 1
        if len({record["중학교_졸업자수"] for record in group_records}) != 1:
            inconsistent_graduate_groups += 1
        by_type = {str(record["고등학교_유형"]): record for record in group_records}
        if set(by_type) == set(SCHOOL_TYPES):
            total = by_type["전체"]["고등학교_진학자수"]
            known_detail_sum = sum(
                int(by_type[school_type]["고등학교_진학자수"] or 0)
                for school_type in SCHOOL_TYPES[1:]
            )
            if total != known_detail_sum:
                detail_total_mismatches += 1

    checked_rate_rows = 0
    rate_formula_mismatches = 0
    for record in records:
        graduates = record["중학교_졸업자수"]
        entrants = record["고등학교_진학자수"]
        official_rate = record["졸업자대비_고등학교진학률_pct"]
        if graduates in (None, 0) or entrants is None or official_rate is None:
            continue
        checked_rate_rows += 1
        if abs(int(entrants) / int(graduates) * 100 - float(official_rate)) > 0.051:
            rate_formula_mismatches += 1

    seoul_total_mismatches = 0
    for year in range(2011, 2026):
        for school_type in SCHOOL_TYPES:
            seoul = next(
                record for record in records
                if record["연도"] == year
                and record["지역"] == "서울특별시"
                and record["고등학교_유형"] == school_type
            )
            district_rows = [
                record for record in records
                if record["연도"] == year
                and record["지역구분"] == "자치구"
                and record["고등학교_유형"] == school_type
            ]
            if sum(int(record["중학교_졸업자수"]) for record in district_rows) != seoul["중학교_졸업자수"]:
                seoul_total_mismatches += 1
            if sum(int(record["고등학교_진학자수"] or 0) for record in district_rows) != seoul["고등학교_진학자수"]:
                seoul_total_mismatches += 1

    record_lookup = {
        (int(record["연도"]), str(record["지역"]), str(record["고등학교_유형"])): record
        for record in records
    }
    for (year, region), group_records in grouped.items():
        by_type = {str(record["고등학교_유형"]): record for record in group_records}
        total = by_type["전체"]
        total_entrants = total["고등학교_진학자수"]
        graduates = total["중학교_졸업자수"]
        for school_type, record in by_type.items():
            if school_type == "전체":
                record["미진학자수"] = (
                    None
                    if graduates is None or total_entrants is None
                    else int(graduates) - int(total_entrants)
                )
                record["유형별_진학구성비_pct"] = None
            else:
                entrants = record["고등학교_진학자수"]
                record["미진학자수"] = None
                record["유형별_진학구성비_pct"] = (
                    None
                    if entrants is None or total_entrants in (None, 0)
                    else round(int(entrants) / int(total_entrants) * 100, 2)
                )

    for record in records:
        year = int(record["연도"])
        region = str(record["지역"])
        school_type = str(record["고등학교_유형"])
        previous = record_lookup.get((year - 1, region, school_type))
        if school_type == "전체":
            record["졸업자수_전년대비증감률_pct"] = (
                None
                if previous is None
                else percent_change(
                    int(record["중학교_졸업자수"]),
                    int(previous["중학교_졸업자수"]),
                )
            )
            record["진학구성비_전년대비증감_pp"] = None
        else:
            record["졸업자수_전년대비증감률_pct"] = None
            current_share = record["유형별_진학구성비_pct"]
            previous_share = previous["유형별_진학구성비_pct"] if previous else None
            record["진학구성비_전년대비증감_pp"] = (
                None
                if current_share is None or previous_share is None
                else round(float(current_share) - float(previous_share), 2)
            )

    records.sort(key=lambda record: (
        int(record["연도"]),
        0 if record["지역구분"] == "서울전체" else 1,
        str(record["지역"]),
        SCHOOL_TYPES.index(str(record["고등학교_유형"])),
    ))

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "진학현황_상세"
    data_sheet.append([COLUMN_LABELS[field] for field in OUTPUT_FIELDS])
    for record in records:
        data_sheet.append([record[field] for field in OUTPUT_FIELDS])
    style_sheet(data_sheet, {
        "기준연도": 10, "지역구분": 12, "지역명": 14, "진학_고등학교유형": 18,
        "중학교_졸업자수(명)": 20, "해당유형_진학자수(명)": 22,
        "졸업자중_해당유형진학률(%)": 30,
        "고등학교_미진학자수(명)": 24, "진학자중_해당유형비율(%)": 26,
        "중학교졸업자수_전년대비증감률(%)": 34,
        "해당유형진학비율_전년대비증감(%p)": 36,
    })

    missing_entrants = sum(record["고등학교_진학자수"] is None for record in records)
    missing_rates = sum(record["졸업자대비_고등학교진학률_pct"] is None for record in records)
    validation_sheet = workbook.create_sheet("검증결과")
    validation_sheet.append(["검사항목", "결과", "설명"])
    validations = [
        ("원본 행 수", len(source_records), "15년 × 26지역 × 6유형 = 2,340행"),
        ("연도+지역+유형 중복", duplicate_count, "0건이어야 함"),
        ("유형 구성 불완전 그룹", incomplete_type_groups, "각 연도·지역에 6개 유형 필요"),
        ("유형별 졸업자수 불일치 그룹", inconsistent_graduate_groups, "같은 연도·지역의 졸업자수는 동일해야 함"),
        ("유형별 진학자 합계 불일치 그룹", detail_total_mismatches, "결측은 결과에서 유지하고 검증 합계에서만 0으로 간주"),
        ("서울전체 합계 불일치", seoul_total_mismatches, "25개 자치구 합과 서울전체 비교"),
        ("진학률 산식 검증 행", checked_rate_rows, "진학자수와 진학률이 모두 있는 행"),
        ("진학률 산식 불일치", rate_formula_mismatches, "소수 첫째 자리 반올림 오차 0.051%p 허용"),
        ("진학자수 결측", missing_entrants, "원본 결측 유지"),
        ("진학률 결측", missing_rates, "원본 결측 유지"),
    ]
    for validation in validations:
        validation_sheet.append(validation)
    style_sheet(validation_sheet, {"검사항목": 34, "결과": 16, "설명": 82})

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_FILE.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(OUTPUT_FILE)
    print(
        f"완료: {len(records):,}행, {len(OUTPUT_FIELDS):,}열, "
        f"진학자수 결측 {missing_entrants}건, 진학률 결측 {missing_rates}건, {OUTPUT_FILE}"
    )


if __name__ == "__main__":
    main()
