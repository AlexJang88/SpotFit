from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE_ROOT = ROOT / "data" / "raw" / "학교" / "학교별_일람표"
OUTPUT_FILE = ROOT / "data" / "interim" / "학교별_일람표_2021_2026_상하반기_통합.csv"


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def find_header_row(worksheet) -> int:
    for row_index in range(1, min(worksheet.max_row, 40) + 1):
        values = {clean(worksheet.cell(row_index, column).value) for column in range(1, worksheet.max_column + 1)}
        if "조사기준일" in values and "학교명" in values:
            return row_index
    raise ValueError(f"헤더 행을 찾지 못했습니다: {worksheet.title}")


def find_header_start(worksheet, header_row: int) -> int:
    for row_index in range(header_row, 0, -1):
        if "학교 기본 정보" in clean(worksheet.cell(row_index, 1).value):
            return row_index
    return header_row


def expanded_header_grid(worksheet, start_row: int, end_row: int) -> list[list[str]]:
    grid = [
        [clean(worksheet.cell(row, column).value) for column in range(1, worksheet.max_column + 1)]
        for row in range(start_row, end_row + 1)
    ]
    for merged in worksheet.merged_cells.ranges:
        if merged.max_row < start_row or merged.min_row > end_row:
            continue
        value = clean(worksheet.cell(merged.min_row, merged.min_col).value)
        for row in range(max(merged.min_row, start_row), min(merged.max_row, end_row) + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                grid[row - start_row][column - 1] = value
    return grid


def unique_headers(worksheet, start_row: int, end_row: int) -> list[str]:
    grid = expanded_header_grid(worksheet, start_row, end_row)
    headers: list[str] = []
    seen: Counter[str] = Counter()
    for column in range(worksheet.max_column):
        levels: list[str] = []
        for row in grid:
            value = row[column]
            if value and (not levels or value != levels[-1]):
                levels.append(value)
        base = "__".join(levels) if levels else f"이름없는열_{column + 1}"
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def find_field(headers: list[str], leaf_name: str) -> int:
    for index, header in enumerate(headers):
        if header.split("__")[-1] == leaf_name:
            return index
    raise ValueError(f"필수 열을 찾지 못했습니다: {leaf_name}")


def main() -> None:
    source_files = sorted(SOURCE_ROOT.rglob("*.xlsx"))
    if len(source_files) != 11:
        raise FileNotFoundError(f"학교별 일람표 XLSX 11개가 필요합니다. 현재 {len(source_files)}개")

    records: list[dict[str, str]] = []
    union_fields: list[str] = ["기준연도", "기준시점", "원본파일"]
    counts: Counter[tuple[int, str]] = Counter()

    for path in source_files:
        period = path.parent.name
        workbook = load_workbook(path, read_only=False, data_only=True)
        worksheet = workbook.worksheets[0]
        header_row = find_header_row(worksheet)
        header_start = find_header_start(worksheet, header_row)
        headers = unique_headers(worksheet, header_start, header_row)
        date_index = find_field(headers, "조사기준일")
        school_index = find_field(headers, "학교명")

        for header in headers:
            if header not in union_fields:
                union_fields.append(header)

        file_count = 0
        for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            date_value = clean(values[date_index] if date_index < len(values) else None)
            school_name = clean(values[school_index] if school_index < len(values) else None)
            date_digits = re.sub(r"\D", "", date_value)
            if len(date_digits) != 8 or not school_name:
                continue
            year = int(date_digits[:4])
            record = {
                header: clean(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
            }
            record.update({"기준연도": str(year), "기준시점": period, "원본파일": path.name})
            records.append(record)
            counts[(year, period)] += 1
            file_count += 1
        print(f"{path.name}: {file_count:,}행")

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_FILE.with_suffix(".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=union_fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
    temporary.replace(OUTPUT_FILE)

    print(f"통합 완료: {len(records):,}행, {len(union_fields):,}열")
    for (year, period), count in sorted(counts.items()):
        print(f"{year} {period}: {count:,}행")
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
