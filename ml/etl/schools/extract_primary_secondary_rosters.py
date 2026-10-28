from __future__ import annotations

import re
import sys
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from ml.etl.schools.extract_kindergarten_rosters import add_admin_dongs


SOURCE_FILE = ROOT / "data" / "interim" / "학교데이터_통합.xlsx"
OUTPUT_FILE = ROOT / "data" / "interim" / "초중고_개별현황_2021_2026_상하반기.xlsx"
TARGET_GRADES = ("초등학교", "중학교", "고등학교")
OFFICIAL_SCHOOL_DONG_OVERRIDES = {
    ("서대문구", "인왕중학교"): "홍제3동",
    ("중랑구", "혜원여자중학교"): "망우제3동",
}

COMMON_FIELDS = {
    "기준연도": ("기준연도",),
    "기준시점": ("기준시점",),
    "원본파일": ("원본파일",),
    "조사기준일": ("학교 기본 정보__조사기준일",),
    "KEDI학교코드": (
        "학교 기본 정보__KEDI학교코드",
        "학교 기본 정보__학교코드 (KEDI)",
    ),
    "학교명": ("학교 기본 정보__학교명",),
    "주소": ("학교 기본 정보__주소",),
    "자치구": ("학교 기본 정보__행정구",),
    "설립구분": ("학교 기본 정보__설립",),
    "학교상태": ("학교 기본 정보__상태",),
    "본분교": ("학교 기본 정보__본분교",),
    "학급수": ("학년별 편성 학급수 (개)__편성학급수_계",),
    "학생수_계": ("학년별 학생수 (명)__총계__학생수_총계_계",),
    "특수학급수": (
        "학급 구분별 학급수 및 학생수__특수학급(편성학급)__특수학급_학급수",
    ),
    "특수학생수_계": (
        "학급 구분별 학급수 및 학생수__특수학급(편성학급)__특수학급_학생수_계",
    ),
    "교원수_계": (
        "교원수 (명)__총계__교원수_총계_계",
        "교원수 (명)__총계__교원수_총계__교원수_총계_계",
    ),
    "특수교사수_계": (
        "교원수 (명)__특수교사(정규)__교원수_정규_특수_계",
        "교원수 (명)__정규교원__교원수_정규교원_특수교사__교원수_정규_특수_계",
    ),
    "학급당_학생수": ("학급당 학생수 (명)__학급당 학생수",),
    "교원1인당_학생수": ("교원1인당 학생수 (명)__교원1인당 학생수",),
}

HIGH_FIELDS = {
    "고등학교유형": ("학교 기본 정보__고등학교 유형",),
    "학교세부유형": (
        "학교 기본 정보__학교세부유형",
        "학교 기본 정보__학교 세부유형",
    ),
    "주야": ("학교 기본 정보__주야",),
    "남녀공학구분": ("학교 기본 정보__남녀공학 구분",),
    "입학자수_계": ("입학 및 졸업 (명)__입학자_계",),
    "졸업자수_계": ("입학 및 졸업 (명)__졸업자_계",),
    "진학자수_계": (
        "졸업 후 상황 (명,%)__진학자_계",
        "졸업 후 상황 (명,%) 고등학교 졸업 후 상황은 12월 제공__진학자_계",
        "졸업 후 상황 (명,%) (고등학교 졸업 후 상황은 11월 제공)__진학자_계",
    ),
    "진학률": (
        "졸업 후 상황 (명,%)__진학률 _전체(%)",
        "졸업 후 상황 (명,%) 고등학교 졸업 후 상황은 12월 제공__진학률 _전체(%)",
        "졸업 후 상황 (명,%) (고등학교 졸업 후 상황은 11월 제공)__진학률 _전체(%)",
    ),
    "취업자수_계": (
        "졸업 후 상황 (명,%)__취업자_계",
        "졸업 후 상황 (명,%) 고등학교 졸업 후 상황은 12월 제공__취업자_계",
        "졸업 후 상황 (명,%) (고등학교 졸업 후 상황은 11월 제공)__취업자_계",
    ),
}


def grade_fields(max_grade: int) -> dict[str, tuple[str, ...]]:
    fields: dict[str, tuple[str, ...]] = {}
    for grade in range(1, max_grade + 1):
        fields[f"{grade}학년_학급수"] = (
            f"학년별 편성 학급수 (개)__{grade}학년_학급수",
        )
        fields[f"{grade}학년_학생수"] = (
            f"학년별 학생수 (명)__{grade}학년__{grade}학년_학생수_계",
        )
    return fields


FIELDS_BY_GRADE = {
    "초등학교": COMMON_FIELDS | grade_fields(6),
    "중학교": COMMON_FIELDS | grade_fields(3),
    "고등학교": COMMON_FIELDS | HIGH_FIELDS | grade_fields(3),
}


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def coalesce(row: tuple[object, ...], indexes: list[int], field: str) -> str:
    values = [clean(row[index]) for index in indexes if clean(row[index])]
    distinct = list(dict.fromkeys(values))
    if len(distinct) > 1:
        raise ValueError(f"{field}에 서로 다른 공식 원본값이 중복 존재합니다: {distinct}")
    return distinct[0] if distinct else ""


def normalize(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]", "", value).lower()


def number_value(value: str) -> object:
    if not value:
        return ""
    try:
        number = Decimal(value.replace(",", ""))
    except InvalidOperation:
        return value
    return int(number) if number == number.to_integral_value() else float(number)


def output_fields(grade: str) -> list[str]:
    fields = list(FIELDS_BY_GRADE[grade])
    fields.insert(fields.index("자치구") + 1, "행정동")
    return fields


def duplicate_counts(records: list[dict[str, str]]) -> tuple[int, int]:
    kedi = Counter(
        (record["기준연도"], record["KEDI학교코드"])
        for record in records
        if record["KEDI학교코드"]
    )
    identity = Counter(
        (
            record["기준연도"],
            record["자치구"],
            normalize(record["학교명"]),
            normalize(record["주소"]),
        )
        for record in records
    )
    return (
        sum(count - 1 for count in kedi.values() if count > 1),
        sum(count - 1 for count in identity.values() if count > 1),
    )


def optimize_sheet(worksheet, fields: list[str]) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 24
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = {
        "기준연도": 10,
        "기준시점": 10,
        "원본파일": 34,
        "조사기준일": 13,
        "KEDI학교코드": 16,
        "학교명": 30,
        "주소": 62,
        "자치구": 12,
        "행정동": 14,
        "설립구분": 11,
        "학교상태": 13,
        "본분교": 10,
        "고등학교유형": 18,
        "학교세부유형": 18,
        "남녀공학구분": 14,
    }
    for index, field in enumerate(fields, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = widths.get(field, 15)


def main() -> None:
    source = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    worksheet = source["학교별_일람표"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows)]
    header_indexes = {header: index for index, header in enumerate(headers)}
    grade_index = next(
        index for index, header in enumerate(headers) if header.split("__")[-1] == "학교급"
    )

    needed = {
        source_field
        for fields in FIELDS_BY_GRADE.values()
        for sources in fields.values()
        for source_field in sources
    }
    missing = sorted(needed - set(headers))
    if missing:
        raise KeyError(f"학교별 일람표에 필요한 컬럼이 없습니다: {missing}")

    indexes_by_grade = {
        grade: {
            output: [header_indexes[source_field] for source_field in sources]
            for output, sources in fields.items()
        }
        for grade, fields in FIELDS_BY_GRADE.items()
    }
    records_by_grade: dict[str, list[dict[str, str]]] = {
        grade: [] for grade in TARGET_GRADES
    }
    for row in rows:
        grade = clean(row[grade_index])
        if grade not in records_by_grade:
            continue
        records_by_grade[grade].append(
            {
                field: coalesce(row, indexes, field)
                for field, indexes in indexes_by_grade[grade].items()
            }
        )
    source.close()

    all_records = [record for grade in TARGET_GRADES for record in records_by_grade[grade]]
    add_admin_dongs(all_records)
    for record in all_records:
        if not record["행정동"]:
            record["행정동"] = OFFICIAL_SCHOOL_DONG_OVERRIDES.get(
                (record["자치구"], record["학교명"]), ""
            )

    workbook = Workbook()
    workbook.remove(workbook.active)
    results: list[str] = []
    for grade in TARGET_GRADES:
        fields = output_fields(grade)
        for point, suffix in (("상반기", "상반기_분석"), ("하반기", "하반기_검증")):
            records = [
                record for record in records_by_grade[grade] if record["기준시점"] == point
            ]
            records.sort(
                key=lambda record: (
                    int(record["기준연도"]),
                    record["자치구"],
                    record["행정동"],
                    record["학교명"],
                )
            )
            duplicates = duplicate_counts(records)
            short_grade = grade.replace("학교", "")
            sheet = workbook.create_sheet(f"{short_grade}_{suffix}")
            sheet.append(fields)
            for record in records:
                sheet.append(
                    [
                        number_value(record.get(field, ""))
                        if field not in {
                            "기준시점",
                            "원본파일",
                            "조사기준일",
                            "KEDI학교코드",
                            "학교명",
                            "주소",
                            "자치구",
                            "행정동",
                            "설립구분",
                            "학교상태",
                            "본분교",
                            "고등학교유형",
                            "학교세부유형",
                            "주야",
                            "남녀공학구분",
                        }
                        else record.get(field, "")
                        for field in fields
                    ]
                )
            optimize_sheet(sheet, fields)
            missing_dong = sum(not record["행정동"] for record in records)
            results.append(
                f"{sheet.title}: {len(records):,}행, {len(fields):,}열, "
                f"행정동 미확인 {missing_dong:,}, 중복(KEDI/학교명+주소) "
                f"{duplicates[0]:,}/{duplicates[1]:,}"
            )

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_FILE.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(OUTPUT_FILE)
    print(f"완료: {OUTPUT_FILE}")
    for result in results:
        print(result)


if __name__ == "__main__":
    main()
