from __future__ import annotations

import csv
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
SOURCE_FILE = ROOT / "data" / "raw" / "범죄" / "범죄_2014_2025.csv"
OUTPUT_FILE = ROOT / "data" / "interim" / "범죄_2014_2025_전처리.xlsx"

COUNT_FIELDS = [
    "소계_발생건",
    "소계_검거건",
    "살인_발생건",
    "살인_검거건",
    "강도_발생건",
    "강도_검거건",
    "강간·강제추행_발생건",
    "강간·강제추행_검거건",
    "절도_발생건",
    "절도_검거건",
    "폭력_발생건",
    "폭력_검거건",
]
OCCURRENCE_FIELDS = [
    "살인_발생건",
    "강도_발생건",
    "강간·강제추행_발생건",
    "절도_발생건",
    "폭력_발생건",
]
ARREST_FIELDS = [
    "살인_검거건",
    "강도_검거건",
    "강간·강제추행_검거건",
    "절도_검거건",
    "폭력_검거건",
]


def parse_count(value: str) -> int | None:
    cleaned = value.strip().replace(",", "")
    if not cleaned:
        return None
    try:
        number = Decimal(cleaned)
    except InvalidOperation as error:
        raise ValueError(f"정수가 아닌 범죄 건수입니다: {value!r}") from error
    if number != number.to_integral_value() or number < 0:
        raise ValueError(f"0 이상의 정수가 아닌 범죄 건수입니다: {value!r}")
    return int(number)


def fill_single_zero(record: dict[str, object], detail_fields: list[str], total_field: str) -> int:
    missing = [field for field in detail_fields if record[field] is None]
    if not missing:
        return 0
    if len(missing) != 1:
        raise ValueError(
            f"{record['연도']} {record['자치구']} {total_field}: "
            f"결측 세부항목이 {len(missing)}개라 산술 복원이 불가능합니다."
        )
    known_sum = sum(int(record[field]) for field in detail_fields if record[field] is not None)
    residual = int(record[total_field]) - known_sum
    if residual != 0:
        raise ValueError(
            f"{record['연도']} {record['자치구']} {missing[0]}: "
            f"소계 차감값이 0이 아닙니다({residual})."
        )
    record[missing[0]] = 0
    return 1


def validate_totals(record: dict[str, object]) -> None:
    occurrence_sum = sum(int(record[field]) for field in OCCURRENCE_FIELDS)
    arrest_sum = sum(int(record[field]) for field in ARREST_FIELDS)
    if occurrence_sum != record["소계_발생건"]:
        raise ValueError(f"{record['연도']} {record['자치구']}: 발생 소계 불일치")
    if arrest_sum != record["소계_검거건"]:
        raise ValueError(f"{record['연도']} {record['자치구']}: 검거 소계 불일치")


def rate_field(field: str) -> str:
    return f"{field}_전년대비증감률(%)"


def main() -> None:
    with SOURCE_FILE.open(encoding="utf-8-sig", newline="") as handle:
        source_records = list(csv.DictReader(handle))

    records: list[dict[str, object]] = []
    filled_count = 0
    for source in source_records:
        record: dict[str, object] = {
            "연도": int(source["연도"]),
            "자치구": source["자치구"].strip(),
        }
        for field in COUNT_FIELDS:
            record[field] = parse_count(source[field])
        filled_count += fill_single_zero(record, OCCURRENCE_FIELDS, "소계_발생건")
        filled_count += fill_single_zero(record, ARREST_FIELDS, "소계_검거건")
        validate_totals(record)
        records.append(record)

    keys = [(record["연도"], record["자치구"]) for record in records]
    if len(keys) != len(set(keys)):
        raise ValueError("연도+자치구 중복 행이 있습니다.")

    years_by_gu: dict[str, list[int]] = defaultdict(list)
    for record in records:
        years_by_gu[str(record["자치구"])].append(int(record["연도"]))
    expected_years = list(range(2014, 2026))
    invalid_years = {
        gu: sorted(years) for gu, years in years_by_gu.items() if sorted(years) != expected_years
    }
    if invalid_years:
        raise ValueError(f"자치구별 연도 구성이 불완전합니다: {invalid_years}")

    records.sort(key=lambda record: (int(record["연도"]), str(record["자치구"])))
    previous_by_gu: dict[str, dict[str, object]] = {}
    for record in records:
        gu = str(record["자치구"])
        previous = previous_by_gu.get(gu)
        for field in COUNT_FIELDS:
            previous_value = int(previous[field]) if previous else 0
            if previous is None or previous_value == 0:
                record[rate_field(field)] = None
            else:
                current_value = int(record[field])
                record[rate_field(field)] = round(
                    (current_value - previous_value) / previous_value * 100,
                    2,
                )
        previous_by_gu[gu] = record

    output_fields = ["연도", "자치구"]
    for field in COUNT_FIELDS:
        output_fields.extend([field, rate_field(field)])

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "범죄_가로형"
    worksheet.append(output_fields)
    for record in records:
        worksheet.append([record.get(field) for field in output_fields])

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = True
    worksheet.row_dimensions[1].height = 30
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, field in enumerate(output_fields, 1):
        width = 10 if field == "연도" else 12 if field == "자치구" else 18
        if "증감률" in field:
            width = 25
            for cell in worksheet[get_column_letter(index)][1:]:
                cell.number_format = "0.00"
        worksheet.column_dimensions[get_column_letter(index)].width = width

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_FILE.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(OUTPUT_FILE)
    print(
        f"완료: {len(records):,}행, {len(output_fields):,}열, "
        f"산술 검증 후 0 복원 {filled_count:,}셀, {OUTPUT_FILE}"
    )


if __name__ == "__main__":
    main()
