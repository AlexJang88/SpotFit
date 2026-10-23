"""어린이 보호구역 공식 통계와 상세 현황을 행정동·기준일 단위로 결합한다."""

from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
INTERIM_DIR = ROOT / "data" / "interim"
OUTPUT_DIR = ROOT / "data" / "processed"
STATS_PATH = INTERIM_DIR / "서울시_행정동별_어린이보호구역_지정통계_통합.csv"
DETAIL_PATH = INTERIM_DIR / "어린이보호구역_지정현황_통합.csv"
FINAL_PATH = OUTPUT_DIR / "행정동별_어린이보호구역_최종.csv"
VALIDATION_PATH = OUTPUT_DIR / "어린이보호구역_집계검증.csv"
UNMATCHED_PATH = OUTPUT_DIR / "어린이보호구역_결합실패.csv"
DEFAULT_EXCEL_PATH = OUTPUT_DIR / "어린이보호구역_결합결과.xlsx"

CATEGORIES = ["초등학교", "유치원", "어린이집", "특수학교", "학원", "외국인학교", "대안학교"]
COMMON_DATES = {"2023-12-31", "2024-06-30", "2024-12-31", "2025-12-31", "2026-06-30"}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def normalize_region(value: str) -> str:
    value = (value or "").strip()
    value = value.replace("·", ".").replace("ㆍ", ".")
    value = re.sub(r"\s+", "", value)
    value = re.sub(r"\.+", ".", value)
    return value


def normalize_facility_type(value: str) -> str:
    value = (value or "").strip()
    if "유치원" in value:
        return "유치원"
    return value


def integer(value: str) -> int:
    cleaned = str(value or "0").replace(",", "").strip()
    return int(float(cleaned)) if cleaned else 0


def write_csv(path: Path, rows: list[dict[str, object]], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def build_outputs():
    stats_rows = [
        row
        for row in read_csv(STATS_PATH)
        if row["기준일"] in COMMON_DATES and row["행정동"] not in {"합계", "소계"}
    ]
    detail_rows = [row for row in read_csv(DETAIL_PATH) if row["기준일"] in COMMON_DATES]

    detail_groups: dict[tuple[str, str, str], dict[str, object]] = {}
    for row in detail_rows:
        key = (
            row["기준일"],
            normalize_region(row["자치구"]),
            normalize_region(row["행정동"]),
        )
        group = detail_groups.setdefault(
            key,
            {
                "기준일": row["기준일"],
                "자치구": row["자치구"].strip(),
                "행정동": row["행정동"].strip(),
                "counts": defaultdict(int),
                "facilities": set(),
                "years": [],
                "unmapped_types": set(),
            },
        )
        facility_type = normalize_facility_type(row["시설유형"])
        if facility_type in CATEGORIES:
            group["counts"][facility_type] += 1
        else:
            group["unmapped_types"].add(row["시설유형"])
        if row["시설명"].strip():
            group["facilities"].add(row["시설명"].strip())
        if row["지정연도"].strip().isdigit():
            group["years"].append(int(row["지정연도"]))

    final_rows: list[dict[str, object]] = []
    validation_rows: list[dict[str, object]] = []
    unmatched_rows: list[dict[str, object]] = []
    matched_keys: set[tuple[str, str, str]] = set()

    for stat in stats_rows:
        key = (
            stat["기준일"],
            normalize_region(stat["자치구"]),
            normalize_region(stat["행정동"]),
        )
        detail = detail_groups.get(key)
        if detail is None:
            unmatched_rows.append(
                {
                    "결합실패유형": "통계에만 존재",
                    "기준일": stat["기준일"],
                    "자치구": stat["자치구"],
                    "행정동": stat["행정동"],
                    "상세_계": "",
                    "미분류_시설유형": "",
                }
            )
            detail_counts = defaultdict(int)
            facilities: list[str] = []
            years: list[int] = []
            unmapped_types: list[str] = []
            join_status = "통계에만 존재"
        else:
            matched_keys.add(key)
            detail_counts = detail["counts"]
            facilities = sorted(detail["facilities"])
            years = detail["years"]
            unmapped_types = sorted(detail["unmapped_types"])
            join_status = "결합완료"

        official_total = integer(stat["계"])
        detail_total = sum(detail_counts[category] for category in CATEGORIES)
        result: dict[str, object] = {
            "기준일": stat["기준일"],
            "자치구": stat["자치구"],
            "행정동": stat["행정동"],
            "결합상태": join_status,
            "통계_계": official_total,
            "상세_계": detail_total,
            "계_차이": official_total - detail_total,
        }
        for category in CATEGORIES:
            official = integer(stat[category])
            detailed = detail_counts[category]
            result[f"통계_{category}"] = official
            result[f"상세_{category}"] = detailed
            result[f"{category}_차이"] = official - detailed
        result.update(
            {
                "시설명목록": " | ".join(facilities),
                "최초지정연도": min(years) if years else "",
                "최근지정연도": max(years) if years else "",
                "미분류_시설유형": " | ".join(unmapped_types),
            }
        )
        final_rows.append(result)

        validation_rows.append(
            {
                "기준일": stat["기준일"],
                "자치구": stat["자치구"],
                "행정동": stat["행정동"],
                "결합상태": join_status,
                "통계_계": official_total,
                "상세_계": detail_total,
                "계_차이": official_total - detail_total,
                "유형별_완전일치": "Y"
                if all(result[f"{category}_차이"] == 0 for category in CATEGORIES)
                else "N",
                "미분류_시설유형": " | ".join(unmapped_types),
            }
        )

    for key, detail in detail_groups.items():
        if key in matched_keys:
            continue
        detail_total = sum(detail["counts"][category] for category in CATEGORIES)
        unmatched_rows.append(
            {
                "결합실패유형": "상세에만 존재",
                "기준일": detail["기준일"],
                "자치구": detail["자치구"],
                "행정동": detail["행정동"],
                "상세_계": detail_total,
                "미분류_시설유형": " | ".join(sorted(detail["unmapped_types"])),
            }
        )

    final_rows.sort(key=lambda row: (row["기준일"], row["자치구"], row["행정동"]))
    validation_rows.sort(key=lambda row: (row["기준일"], row["자치구"], row["행정동"]))
    unmatched_rows.sort(key=lambda row: (row["기준일"], row["결합실패유형"], row["자치구"], row["행정동"]))
    return final_rows, validation_rows, unmatched_rows


def make_excel(path: Path, datasets: list[tuple[str, list[dict[str, object]], list[str]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for title, rows, columns in datasets:
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in rows:
            sheet.append([row.get(column, "") for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(columns, start=1):
            width = 18 if column != "시설명목록" else 60
            sheet.column_dimensions[get_column_letter(index)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel-output", type=Path, default=DEFAULT_EXCEL_PATH)
    args = parser.parse_args()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    final_rows, validation_rows, unmatched_rows = build_outputs()
    final_columns = list(final_rows[0])
    validation_columns = list(validation_rows[0])
    unmatched_columns = list(unmatched_rows[0])
    write_csv(FINAL_PATH, final_rows, final_columns)
    write_csv(VALIDATION_PATH, validation_rows, validation_columns)
    write_csv(UNMATCHED_PATH, unmatched_rows, unmatched_columns)
    make_excel(
        args.excel_output,
        [
            ("행정동별 최종", final_rows, final_columns),
            ("집계검증", validation_rows, validation_columns),
            ("결합실패", unmatched_rows, unmatched_columns),
        ],
    )
    exact = sum(row["유형별_완전일치"] == "Y" for row in validation_rows)
    print(f"최종 {len(final_rows):,}행, 완전일치 {exact:,}행, 결합실패 {len(unmatched_rows):,}행")
    print(args.excel_output.resolve())


if __name__ == "__main__":
    main()
