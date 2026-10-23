"""통계·상세 결합본을 검증용으로 보관하고 공식 통계 중심 최종본을 만든다."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
INTERIM_DIR = ROOT / "data" / "interim"
PROCESSED_DIR = ROOT / "data" / "processed"
CURRENT_FINAL = PROCESSED_DIR / "행정동별_어린이보호구역_최종.csv"
VALIDATION_COPY = INTERIM_DIR / "행정동별_어린이보호구역_통계상세_검증용.csv"
VALIDATION_EXCEL = INTERIM_DIR / "행정동별_어린이보호구역_통계상세_검증용.xlsx"
FINAL_CSV = PROCESSED_DIR / "행정동별_어린이보호구역_최종.csv"
FINAL_EXCEL = PROCESSED_DIR / "행정동별_어린이보호구역_최종.xlsx"

COLUMN_MAP = {
    "기준일": "기준일",
    "자치구": "자치구",
    "행정동": "행정동",
    "통계_계": "계",
    "통계_초등학교": "초등학교",
    "통계_유치원": "유치원",
    "통계_어린이집": "어린이집",
    "통계_특수학교": "특수학교",
    "통계_학원": "학원",
    "통계_외국인학교": "외국인학교",
    "통계_대안학교": "대안학교",
}


def write_excel(csv_path: Path, excel_path: Path, sheet_name: str) -> None:
    with csv_path.open(encoding="utf-8-sig", newline="") as source:
        rows = list(csv.reader(source))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(rows[0], start=1):
        width = 60 if header == "시설명목록" else 18
        sheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(excel_path)


def main() -> None:
    INTERIM_DIR.mkdir(parents=True, exist_ok=True)
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)

    # 상세 재집계와 차이 컬럼이 포함된 기존 32열 결과는 검증용으로 보존한다.
    shutil.copy2(CURRENT_FINAL, VALIDATION_COPY)
    write_excel(VALIDATION_COPY, VALIDATION_EXCEL, "통계상세 검증용")

    with VALIDATION_COPY.open(encoding="utf-8-sig", newline="") as source:
        rows = list(csv.DictReader(source))

    with FINAL_CSV.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=list(COLUMN_MAP.values()))
        writer.writeheader()
        for row in rows:
            writer.writerow({new: row[old] for old, new in COLUMN_MAP.items()})

    write_excel(FINAL_CSV, FINAL_EXCEL, "행정동별 통계")
    print(f"검증용 보관: {VALIDATION_COPY}")
    print(f"최종 통계: {FINAL_CSV}")
    print(f"최종 통계: {FINAL_EXCEL}")
    print(f"행 수: {len(rows):,}, 최종 열 수: {len(COLUMN_MAP)}")


if __name__ == "__main__":
    main()
