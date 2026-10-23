"""어린이 보호구역의 행정동별 통계와 개별 지정현황을 기준일별로 결합한다."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date
from pathlib import Path
from xml.etree import ElementTree

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "data" / "raw" / "아동" / "어린이보호구역"
OUTPUT_DIR = ROOT / "data" / "interim"
STATS_OUTPUT = OUTPUT_DIR / "서울시_행정동별_어린이보호구역_지정통계_통합.csv"
DETAIL_OUTPUT = OUTPUT_DIR / "어린이보호구역_지정현황_통합.csv"

DATE_PATTERN = re.compile(r"\((20\d{2})\.\s*(\d{1,2})월말\s*기준\)")
STATS_COLUMNS = [
    "기준일",
    "자치구",
    "행정동",
    "계",
    "초등학교",
    "유치원",
    "어린이집",
    "특수학교",
    "학원",
    "외국인학교",
    "대안학교",
]
DETAIL_COLUMNS = [
    "기준일",
    "연번",
    "시도",
    "자치구",
    "행정동",
    "도로명 주소",
    "시설명",
    "시설유형",
    "지정연도",
]


def reference_date(path: Path) -> str:
    match = DATE_PATTERN.search(path.name)
    if not match:
        raise ValueError(f"파일명에서 기준일을 찾을 수 없습니다: {path.name}")
    year, month = map(int, match.groups())
    last_day = 30 if month == 6 else 31
    return date(year, month, last_day).isoformat()


def load_workbook_safely(path: Path) -> openpyxl.Workbook:
    """일부 2021년 파일의 누락된 스타일명을 메모리에서만 보완해 읽는다."""
    try:
        return openpyxl.load_workbook(path, data_only=True, read_only=True)
    except TypeError as error:
        if "NamedCellStyle" not in str(error):
            raise

    source = zipfile.ZipFile(path)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as repaired:
        for member in source.namelist():
            content = source.read(member)
            if member == "xl/styles.xml":
                root = ElementTree.fromstring(content)
                namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                for index, style in enumerate(root.iter(f"{namespace}cellStyle"), start=1):
                    if "name" not in style.attrib:
                        style.set("name", f"RecoveredStyle{index}")
                content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            elif member.startswith("xl/worksheets/sheet") and member.endswith(".xml"):
                # 2021년 파일에는 실제 스타일 목록 범위를 벗어난 셀 스타일 번호도 있다.
                # 병합에는 서식이 필요하지 않으므로 셀의 스타일 참조만 제거한다.
                content = re.sub(rb' s="\d+"', b"", content)
            repaired.writestr(member, content)
    source.close()
    buffer.seek(0)
    return openpyxl.load_workbook(buffer, data_only=True, read_only=False)


def merge_stats(files: list[Path]) -> tuple[int, dict[str, int]]:
    counts: dict[str, int] = {}
    total = 0
    with STATS_OUTPUT.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output)
        writer.writerow(STATS_COLUMNS)
        for path in files:
            reference = reference_date(path)
            workbook = load_workbook_safely(path)
            sheet = workbook["1. 지정 통계"]
            count = 0
            for row in sheet.iter_rows(min_row=4, min_col=2, max_col=11, values_only=True):
                if row[0] is None or row[1] is None:
                    continue
                writer.writerow([reference, *row])
                count += 1
            workbook.close()
            counts[reference] = count
            total += count
    return total, counts


def detail_sheet(workbook: openpyxl.Workbook):
    candidates = [
        sheet
        for sheet in workbook.worksheets
        if sheet.title != "Sheet1" and "지정대상" not in sheet.title
    ]
    if len(candidates) != 1:
        raise ValueError(f"상세현황 시트를 특정할 수 없습니다: {[s.title for s in candidates]}")
    return candidates[0]


def merge_details(files: list[Path]) -> tuple[int, dict[str, int]]:
    counts: dict[str, int] = {}
    total = 0
    with DETAIL_OUTPUT.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output)
        writer.writerow(DETAIL_COLUMNS)
        for path in files:
            reference = reference_date(path)
            workbook = load_workbook_safely(path)
            sheet = detail_sheet(workbook)
            count = 0
            for row in sheet.iter_rows(values_only=True):
                serial_index = next(
                    (index for index, value in enumerate(row[:2]) if isinstance(value, int)),
                    None,
                )
                if serial_index is None:
                    continue

                if reference == "2021-06-30":
                    # 2021년은 주소 열 없이 동명, 시설명, 시설종류, 지정일자만 제공한다.
                    serial, city, district, dong, facility, facility_type, year = row[:7]
                    address = None
                    if city == "서울":
                        city = "서울특별시"
                elif serial_index == 0:
                    serial, city, district, dong, address, facility, facility_type, year = row[:8]
                else:
                    serial, district, dong, address, facility, facility_type, year = row[1:8]
                    city = "서울특별시"

                writer.writerow(
                    [reference, serial, city, district, dong, address, facility, facility_type, year]
                )
                count += 1
            workbook.close()
            counts[reference] = count
            total += count
    return total, counts


def main() -> None:
    stats_files = sorted(
        INPUT_DIR.glob("서울시 행정동별 어린이 보호구역 지정 통계*.xlsx"),
        key=reference_date,
    )
    detail_files = sorted(
        INPUT_DIR.glob("어린이 보호구역 지정현황*.xlsx"),
        key=reference_date,
    )
    if len(stats_files) != 5:
        raise RuntimeError(f"지정 통계 파일은 5개여야 합니다: {len(stats_files)}개")
    if len(detail_files) != 9:
        raise RuntimeError(f"지정현황 파일은 9개여야 합니다: {len(detail_files)}개")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stats_total, stats_counts = merge_stats(stats_files)
    detail_total, detail_counts = merge_details(detail_files)

    print(f"지정 통계: {len(stats_files)}개 파일, {stats_total:,}행 {stats_counts}")
    print(STATS_OUTPUT)
    print(f"지정현황: {len(detail_files)}개 파일, {detail_total:,}행 {detail_counts}")
    print(DETAIL_OUTPUT)


if __name__ == "__main__":
    main()
