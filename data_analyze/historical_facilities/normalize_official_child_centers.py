from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE_DIR = ROOT / "data" / "raw" / "5년치_시설" / "지역아동센터"
OUTPUT_DIR = ROOT / "data" / "processed" / "historical_facilities" / "official_2021_2025" / "child_centers"

OUTPUT_FIELDS = [
    "year",
    "source_row",
    "serial",
    "sido",
    "sigungu",
    "center_name",
    "address",
    "phone",
    "fax",
    "homepage",
    "capacity",
    "current_children",
    "source_file",
    "quality_flag",
]

ALIASES = {
    "serial": ("연번",),
    "sido": ("시도", "시도명"),
    "sigungu": ("시군구", "시군구명"),
    "center_name": ("시설명", "센터명"),
    "address": ("주소", "시설 소재지", "소재지"),
    "phone": ("시설 전화번호", "대표 전화번호", "전화번호"),
    "fax": ("팩스 번호", "팩스번호"),
    "homepage": ("홈페이지 주소", "홈페이지"),
    "capacity": ("입소정원", "정원", "정원수"),
    "current_children": ("아동현원", "현원", "현원수"),
}


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_header(rows: list[tuple[object, ...]]) -> int:
    for index, row in enumerate(rows):
        values = {clean(value) for value in row}
        if "연번" in values and ("시설명" in values or "센터명" in values):
            return index
    raise ValueError("연번과 시설명이 포함된 머리글 행을 찾지 못했습니다.")


def canonical_indices(header: tuple[object, ...]) -> dict[str, int]:
    normalized = [clean(value).replace("\n", " ") for value in header]
    result: dict[str, int] = {}
    for canonical, aliases in ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                result[canonical] = normalized.index(alias)
                break
    required = {"serial", "sido", "sigungu", "center_name", "address"}
    missing = sorted(required - result.keys())
    if missing:
        raise ValueError(f"필수 열 누락: {missing}; 원본 머리글={normalized}")
    return result


def normalize_file(path: Path, year: int) -> list[dict[str, str | int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    header_index = find_header(rows)
    indices = canonical_indices(rows[header_index])
    normalized: list[dict[str, str | int]] = []
    for source_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = {key: clean(row[index] if index < len(row) else None) for key, index in indices.items()}
        if not values.get("center_name"):
            continue
        normalized.append(
            {
                "year": year,
                "source_row": source_index,
                **{field: values.get(field, "") for field in OUTPUT_FIELDS if field in ALIASES},
                "source_file": path.name,
                "quality_flag": "official_year_end_address_list",
            }
        )
    return normalized


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    summaries: list[dict[str, int]] = []
    for year in range(2021, 2026):
        matches = sorted(SOURCE_DIR.glob(f"{year}_*.xlsx"))
        if len(matches) != 1:
            raise FileNotFoundError(f"{year}년 원본은 정확히 1개여야 합니다: {matches}")
        rows = normalize_file(matches[0], year)
        output = OUTPUT_DIR / f"{year}.csv"
        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=OUTPUT_FIELDS)
            writer.writeheader()
            writer.writerows(rows)
        seoul = sum(row["sido"] == "서울특별시" for row in rows)
        summaries.append({"year": year, "national_centers": len(rows), "seoul_centers": seoul})

    with (OUTPUT_DIR / "summary.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["year", "national_centers", "seoul_centers"])
        writer.writeheader()
        writer.writerows(summaries)
    print(f"정규화 완료: {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
